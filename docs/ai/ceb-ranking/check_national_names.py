# -*- coding: utf-8 -*-
"""Έλεγχος των λιστών των ομοσπονδιών (docs/ai/ceb-ranking/national) πριν μπουν στη βαθμολογία.

Δύο ερωτήματα:
  1) Ταιριάζουν τα ονόματα που στέλνει η ομοσπονδία με τους παίκτες της βάσης μας (bt-players);
  2) Η σειρά του αρχείου βγάζει τους ίδιους πόντους με τις στήλες B/C/D του επίσημου PDF;

Χρήση:  uv run --python 3.12 --with openpyxl --with xlrd python check_national_names.py
"""
import json, os, re, sys, unicodedata, urllib.request
from collections import defaultdict

import openpyxl
import xlrd

HERE = os.path.dirname(os.path.abspath(__file__))
NAT = os.path.join(HERE, "national")
PUB = os.path.join(HERE, "..", "..", "..", "public", "data", "ceb-ranking", "3c-individual.json")
API = "https://app.billiardtoday.com/api/bt-players"
SCALE = [(1, 1, 40), (2, 2, 27), (3, 4, 19), (5, 8, 13), (9, 16, 8), (17, 32, 4)]
TIE_BREAK_NOTE = "Σε ισοβαθμία η θέση δίνει τους πόντους της ομάδας θέσεων (3-4 -> 19 κ.λπ.), όπως στο PDF."
MAP = str.maketrans({"Ø": "O", "ø": "O", "Æ": "AE", "æ": "AE", "Å": "A", "å": "A", "Ö": "O", "ö": "O",
                     "Ü": "U", "ü": "U", "ß": "SS", "Đ": "D", "đ": "D", "É": "E", "È": "E", "Ê": "E",
                     "Ç": "C", "Ñ": "N", "İ": "I", "ı": "I"})


def norm(s):
    if s is None:
        return ""
    s = str(s).translate(MAP)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    s = re.sub(r"[^0-9A-ZΑ-Ω ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def implied(pos):
    for lo, hi, pts in SCALE:
        if lo <= pos <= hi:
            return pts
    return 0


def fetch_db():
    rows, page = [], 1
    while True:
        url = (API + "?fields%5B0%5D=full_name&fields%5B1%5D=full_name_en&fields%5B2%5D=country"
                     f"&pagination%5BpageSize%5D=1000&pagination%5Bpage%5D={page}")
        with urllib.request.urlopen(url, timeout=60) as r:
            d = json.load(r)
        rows += d["data"]
        if page >= d["meta"]["pagination"]["pageCount"]:
            break
        page += 1
    return [{"id": r["id"], "name": r.get("full_name"), "en": r.get("full_name_en"),
             "cc": (r.get("country") or "").upper(),
             "tok": set((norm(r.get("full_name")) + " " + norm(r.get("full_name_en"))).split())} for r in rows]


# ------------------------------------------------------------------ ανάγνωση αρχείων
def danish(fn):
    ws = openpyxl.load_workbook(os.path.join(NAT, fn), data_only=True).worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], (int, float)) and r[1] and str(r[1]).strip().lower() != "cancellation":
            pid = next((c for c in r[2:] if c not in (None, "")), None)
            out.append((int(r[0]), str(r[1]).strip(), str(pid) if pid else None))
    return out


def greece(block_wanted="LADIES"):
    ws = openpyxl.load_workbook(os.path.join(NAT, "Greece national 3C.xlsx"), data_only=True).worksheets[0]
    on, out = False, []
    for r in ws.iter_rows(values_only=True):
        a, b = (r + (None, None))[:2] if r else (None, None)
        if a is not None and not isinstance(a, (int, float)) and "3C" in str(a).upper():
            on = block_wanted.upper() in str(a).upper()
        elif on and isinstance(a, (int, float)) and b:
            out.append((int(a), str(b).strip(), str(r[2]) if len(r) > 2 and r[2] else None))
    return out


def norway():
    ws = openpyxl.load_workbook(os.path.join(NAT, "UMB Ranking Points 2026 Norway.xlsx"), data_only=True).worksheets[0]
    return [(int(r[0]), str(r[1]).strip(), str(r[2]) if len(r) > 2 and r[2] not in (None, "", 0) else None)
            for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], (int, float)) and r[1]]


def spain_ladies():
    ws = openpyxl.load_workbook(os.path.join(NAT, "General Classification National Championship Ladies 3 Cushions June 2026.xlsx"), data_only=True).worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        if r and len(r) > 5 and isinstance(r[3], (int, float)) and r[5]:
            out.append((int(r[3]), str(r[5]).strip(), str(r[4]) if r[4] not in (None, "") else None))
    return out


def france(section="3 BANDES MASTERS"):
    bk = xlrd.open_workbook(os.path.join(NAT, "Résultats 2025-2026 pour CEB.xls"))
    out = []
    for sh in bk.sheets():
        sec, active = None, False
        for i in range(sh.nrows):
            vals = [sh.cell_value(i, j) for j in range(sh.ncols)]
            first = str(vals[0]).strip() if str(vals[0]).strip() else ""
            if first and not isinstance(vals[0], float) and not any(str(v).strip().upper() == "JOUEURS" for v in vals):
                sec = first
                active = sec.upper() == section.upper()
            elif active and len(vals) > 1 and isinstance(vals[0], float) and str(vals[1]).strip():
                out.append((int(vals[0]), str(vals[1]).strip().title(), None))
    return out


LISTS = [
    ("ΔΑΝΙΑ Άνδρες 3C (2026)", "DK", lambda: danish("Danish Championship 2026 3-Cushion (Mens).xlsx")),
    ("ΔΑΝΙΑ Γυναίκες 3C (2026)", "DK", lambda: danish("Danish Championship 2026 3-Cushion (Ladies).xlsx")),
    ("ΕΛΛΑΔΑ Γυναίκες 3C", "GR", lambda: greece("LADIES")),
    ("ΝΟΡΒΗΓΙΑ 3C (NM 2025)", "NO", norway),
    ("ΙΣΠΑΝΙΑ Γυναίκες 3C", "ES", spain_ladies),
    ("ΓΑΛΛΙΑ 3 BANDES MASTERS", "FR", lambda: france("3 BANDES MASTERS")),
    ("ΓΑΛΛΙΑ Γυναίκες 3 bandes", "FR", lambda: france("3 bandes Dames")),
]


def main():
    db = fetch_db()
    print(f"### βάση μας: {len(db)} παίκτες (bt-players, πεδία full_name/full_name_en — ΧΩΡΙΣ UMB ID)\n")
    ceb = defaultdict(list)
    d = json.load(open(PUB, encoding="utf-8"))
    for r in d["rows"]:
        nat = {k: r["ev"][ord(k) - 65] for k in ("B", "C", "D") if r["ev"][ord(k) - 65]}
        if nat:
            ceb[r["fed"]].append({"name": r["name"], "rank": r["rank"], "nat": nat, "tok": set(norm(r["name"]).split())})
    tot = [0, 0, 0, 0]
    for label, fed, fx in LISTS:
        lst = fx()
        ok, amb, no = [], [], []
        col = None
        cols = defaultdict(int)
        for x in ceb[fed]:
            for k in x["nat"]:
                cols[k] += 1
        if cols:
            col = max(cols, key=cols.get)
        print("=" * 100)
        print(f"### {label}: {len(lst)} παίκτες")
        for pos, name, pid in lst:
            t = set(norm(name).split())
            hits = [x for x in db if t <= x["tok"]]
            if not hits:
                sur = max(t, key=len) if t else ""
                hits = [x for x in db if sur in x["tok"]] if len(sur) >= 4 else []
            same = [h for h in hits if h["cc"] == fed]
            if len(same) == 1:
                hits = same
            (ok if len(hits) == 1 else amb if hits else no).append((pos, name, hits))
        print(f"  προς τη βάση μας: ✓ {len(ok)} μοναδικό | ~ {len(amb)} διφορούμενο | ✗ {len(no)} δεν βρέθηκε")
        for pos, name, hits in amb[:5]:
            print(f"    ~ #{pos} {name!r} -> {[(h['name'] or '') + ' (' + h['cc'] + ', id' + str(h['id']) + ')' for h in hits[:4]]}")
        for pos, name, hits in no[:8]:
            print(f"    ✗ #{pos} {name!r}")
        # έλεγχος πόντων έναντι PDF
        if col and "Γυναίκες" in label:
            print("  έναντι του PDF: δεν ισχύει — η δημοσιευμένη λίστα 3-Cushion είναι ανδρών "
                  "(οι γυναικείες πάνε στη «3-Cushion Ladies»)")
            col = None
        if col:
            good = diff = miss = 0
            for pos, name, pid in lst:
                t = set(norm(name).split())
                cand = [x for x in ceb[fed] if t == x["tok"]] or [x for x in ceb[fed] if t <= x["tok"]]
                if not cand:
                    miss += 1
                    continue
                if cand[0]["nat"].get(col) == implied(pos):
                    good += 1
                else:
                    diff += 1
                    print(f"    ✗ πόντου: #{pos} {name!r} -> PDF {cand[0]['name']} {col}={cand[0]['nat'].get(col)}, το αρχείο δίνει {implied(pos)}")
            print(f"  έναντι του PDF (στήλη {col}): ✓ {good} | διαφορές {diff} | χωρίς αντιστοίχιση {miss}")
            print("  " + TIE_BREAK_NOTE)
        tot = [tot[0] + len(lst), tot[1] + len(ok), tot[2] + len(amb), tot[3] + len(no)]
    n, o, a, x = tot
    print("=" * 100)
    print(f"ΣΥΝΟΛΟ: {n} ονόματα -> μοναδικό {o} ({o/n:.0%}) | διφορούμενο {a} | δεν βρέθηκε {x}")


if __name__ == "__main__":
    sys.exit(main())
