"""Build the Excel file from the verified transcription of the WhatsApp image."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROWS = [
    ("1",     "0223", "David Pennör",         1.399),
    ("2",     "0150", "Torbjörn Blomdahl",    1.970),
    ("3",     "0438", "Henryk Kalita",        0.978),
    ("3",     "3062", "Yüksel Polat",         0.909),
    ("5-8",   "0197", "Nalle Olsson",         1.335),
    ("5-8",   "2914", "Alf Alvberger",        0.904),
    ("5-8",   "0287", "Bosse Gustafsson",     0.810),
    ("5-8",   "2915", "Jesper Johansson",     0.706),
    ("9-16",  "2740", "Ratcha Alievski",      0.737),
    ("9-16",  "3395", "Mio Svensson",         0.705),
    ("9-16",  "2382", "Manuel Orellana",      0.667),
    ("9-16",  "2746", "Ken Hultqvist",        0.615),
    ("9-16",  "2921", "Peter Åkesson",        0.599),
    ("9-16",  "2918", "Mark Jenabpour",       0.558),
    ("9-16",  "2920", "Göran Espe",           0.595),
    ("9-16",  "2928", "Andreas Lagemyr",      0.476),
    ("17-24", "3547", "Adrian Weiss",         0.755),
    ("17-24", "3246", "Mustafa Kasemi",       0.570),
    ("17-24", "2926", "Mats Lohmander",       0.530),
    ("17-24", "2923", "Olle Nisslert",        0.502),
    ("17-24", "3548", "Emad Tavakolizadeh",   0.500),
    ("17-24", "2929", "Uldis Martinovs",      0.490),
    ("17-24", "2748", "Polyzois Arvanitidis", 0.468),
]

TITLE = "Results Swedish Championship 3C 2026"
OUT = r"D:\Projects\4-billiardtoday-frontend\docs\ai\ceb-ranking\sweden-2026-shc-3c-results.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Results 3C 2026"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="DCE6F1")

ws["A1"] = TITLE
ws["A1"].font = Font(bold=True, size=13)
ws.merge_cells("A1:D1")
ws["A1"].alignment = Alignment(horizontal="center")
ws["A2"] = "Svenska Biljardförbundet / Swedish Billiards Federation"
ws["A2"].font = Font(size=9, italic=True, color="666666")
ws.merge_cells("A2:D2")
ws["A2"].alignment = Alignment(horizontal="center")

HEADERS = ["Position", "UMB-ID", "Name", "Average"]
for c, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(bold=True)
    cell.fill = hdr_fill
    cell.border = border
    cell.alignment = Alignment(horizontal="center")

for i, (pos, uid, name, avg) in enumerate(ROWS):
    r = 5 + i
    a = ws.cell(row=r, column=1, value=pos)
    a.alignment = Alignment(horizontal="left")
    b = ws.cell(row=r, column=2, value=uid)
    b.alignment = Alignment(horizontal="center")
    b.number_format = "@"          # keep the zero-padded id as text
    c = ws.cell(row=r, column=3, value=name)
    c.alignment = Alignment(horizontal="left")
    d = ws.cell(row=r, column=4, value=avg)
    d.number_format = "0.000"
    d.alignment = Alignment(horizontal="right")
    for cell in (a, b, c, d):
        cell.border = border

for col, w in zip("ABCD", (11, 10, 26, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# --- provenance sheet ------------------------------------------------------
meta = wb.create_sheet("Meta")
meta_rows = [
    ("Source",            "WhatsApp photo of \"Results Sweden 2026.xlsx\" (Svenska Biljardförbundet)"),
    ("Image",             "docs/ai/ceb-ranking/WhatsApp Image 2026-10-06 at 21.12.46.jpeg"),
    ("Document title",    TITLE),
    ("Kind of list",      "Single event results (Swedish Championship 3C 2026) - NOT a season ranking"),
    ("Rows",              len(ROWS)),
    ("Transcribed",       "2026-10-06 (Hermes)"),
    ("Method",            "tesseract per-cell OCR + independent vision read of upscaled column strips; "
                          "disagreements re-cropped and re-read (row 12 average: 0.615)"),
    ("Note - decimals",   "Source prints Swedish decimals with a comma (1,399). Stored here as numbers."),
    ("Note - order",      "Within the 5-8 / 9-16 / 17-24 groups the order is exactly as printed (not sorted by average)."),
    ("Note - 17-24 group","The source lists 7 players under the position label \"17-24\" (23 rows in total); "
                          "no eighth row is visible anywhere in the screenshot."),
    ("Note - ids",        "UMB-ID stored as text to keep the leading zero (0223)."),
]
meta["A1"] = "Source / verification"
meta["A1"].font = Font(bold=True)
for i, (k, v) in enumerate(meta_rows, start=3):
    meta.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = meta.cell(row=i, column=2, value=v)
    c.alignment = Alignment(wrap_text=True, vertical="top")
meta.column_dimensions["A"].width = 20
meta.column_dimensions["B"].width = 95

wb.save(OUT)
print("saved", OUT, "rows", len(ROWS))
